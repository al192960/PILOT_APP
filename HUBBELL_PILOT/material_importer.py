from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from database import conectar


HOJA_MATERIALES = "MATERIALS"


# ==================================================
# LEER ARCHIVO XLSX
# ==================================================

def leer_materiales_xlsx(ruta_archivo):

    ruta = Path(ruta_archivo)

    if not ruta.exists():

        raise FileNotFoundError(
            "The selected Excel file does not exist."
        )

    # Abrimos el Excel en modo normal.
    # Esto evita problemas con max_row en algunos archivos.
    workbook = load_workbook(
        ruta,
        data_only=True
    )

    try:

        if HOJA_MATERIALES not in workbook.sheetnames:

            raise ValueError(
                "The Excel file must contain "
                "a sheet named MATERIALS."
            )

        hoja = workbook[HOJA_MATERIALES]

        # ------------------------------------------
        # VALIDAR ENCABEZADOS
        # ------------------------------------------

        encabezado_material = hoja["A1"].value
        encabezado_descripcion = hoja["B1"].value

        if (
            str(encabezado_material)
            .strip()
            .upper()
            != "MATERIAL"
        ):

            raise ValueError(
                "Cell A1 must contain MATERIAL."
            )

        if (
            str(encabezado_descripcion)
            .strip()
            .upper()
            != "DESCRIPTION"
        ):

            raise ValueError(
                "Cell B1 must contain DESCRIPTION."
            )

        materiales = {}

        # ------------------------------------------
        # LEER FILAS
        # ------------------------------------------

        for fila in hoja.iter_rows(
            min_row=2,
            min_col=1,
            max_col=2,
            values_only=True
        ):

            valor_material = fila[0]
            valor_descripcion = fila[1]

            if valor_material is None:
                continue

            material = str(
                valor_material
            ).strip()

            if not material:
                continue

            descripcion = ""

            if valor_descripcion is not None:

                descripcion = str(
                    valor_descripcion
                ).strip()

            # Si un material aparece más de una vez
            # en el Excel, se conserva la última fila.

            materiales[material] = descripcion

        if not materiales:

            raise ValueError(
                "No materials were found "
                "in the MATERIALS sheet."
            )

        return materiales

    finally:

        workbook.close()


# ==================================================
# IMPORTAR A SQLITE
# ==================================================

def importar_materiales(
    ruta_archivo,
    modo="ADD_UPDATE"
):

    if modo not in (
        "ADD_UPDATE",
        "REPLACE"
    ):

        raise ValueError(
            "Invalid import mode."
        )

    materiales = leer_materiales_xlsx(
        ruta_archivo
    )

    fecha = (
        datetime
        .now()
        .astimezone()
        .isoformat(
            timespec="seconds"
        )
    )

    conexion = conectar()
    cursor = conexion.cursor()

    agregados = 0
    actualizados = 0
    reactivados = 0
    desactivados = 0

    try:

        cursor.execute(
            "BEGIN IMMEDIATE"
        )

        # ==========================================
        # ADD / UPDATE
        # ==========================================

        for material, descripcion in materiales.items():

            cursor.execute(
                """
                SELECT
                    description,
                    active
                FROM materials
                WHERE material_number = ?
                """,
                (
                    material,
                )
            )

            existente = cursor.fetchone()

            # --------------------------------------
            # MATERIAL NUEVO
            # --------------------------------------

            if existente is None:

                cursor.execute(
                    """
                    INSERT INTO materials (
                        material_number,
                        description,
                        active,
                        created_at,
                        updated_at
                    )
                    VALUES (?, ?, 1, ?, ?)
                    """,
                    (
                        material,
                        descripcion,
                        fecha,
                        fecha
                    )
                )

                agregados += 1

            # --------------------------------------
            # MATERIAL EXISTENTE
            # --------------------------------------

            else:

                descripcion_anterior = (
                    existente[0] or ""
                )

                estaba_activo = (
                    existente[1] == 1
                )

                cursor.execute(
                    """
                    UPDATE materials
                    SET
                        description = ?,
                        active = 1,
                        updated_at = ?
                    WHERE material_number = ?
                    """,
                    (
                        descripcion,
                        fecha,
                        material
                    )
                )

                if not estaba_activo:

                    reactivados += 1

                if (
                    descripcion_anterior
                    != descripcion
                ):

                    actualizados += 1

        # ==========================================
        # REPLACE DATABASE
        # ==========================================

        if modo == "REPLACE":

            materiales_importados = set(
                materiales.keys()
            )

            cursor.execute(
                """
                SELECT material_number
                FROM materials
                WHERE active = 1
                """
            )

            materiales_activos = {
                fila[0]
                for fila in cursor.fetchall()
            }

            materiales_a_desactivar = (
                materiales_activos
                - materiales_importados
            )

            for material in materiales_a_desactivar:

                cursor.execute(
                    """
                    UPDATE materials
                    SET
                        active = 0,
                        updated_at = ?
                    WHERE material_number = ?
                    """,
                    (
                        fecha,
                        material
                    )
                )

                desactivados += 1

        conexion.commit()

    except Exception:

        conexion.rollback()
        raise

    finally:

        conexion.close()

    return {
        "total_excel": len(materiales),
        "added": agregados,
        "updated": actualizados,
        "reactivated": reactivados,
        "deactivated": desactivados
    }