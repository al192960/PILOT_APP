from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from database import conectar


# ==================================================
# CONFIGURACIÓN DE TU PLANTILLA  (EDITA ESTA SECCIÓN)
# ==================================================

# Tu archivo Excel ya preparado. Ponlo en la carpeta del proyecto
# con este nombre, o cambia el nombre aquí.
PLANTILLA_PACKING = (
    Path(__file__).resolve().parent / "packing_template.xlsx"
)

# Nombre de la hoja donde se escribe. None = la hoja activa.
HOJA_PLANTILLA = None

# --------------------------------------------------
# CELDAS FIJAS: un dato que va en una sola celda.
# Formato:  "campo": "CELDA"
#
# Campos disponibles:
#   "date"           fecha y hora de generación
#   "user_name"      nombre del usuario que genera
#   "user_id"        employee_id del usuario que genera
#   "total_pallets"  cantidad de tarimas escaneadas
#   "total_lines"    cantidad de líneas (materiales)
#   "total_qty"      suma de cantidades
#   "destination"    destino de la primera tarima
# --------------------------------------------------
CELDAS_FIJAS = {
    "date": "H1",
    "user_name": "H2",
    "total_pallets": "H3",
    # "destination": "B2",
}

# --------------------------------------------------
# TABLA: una fila por cada línea de cada tarima escaneada,
# empezando en "fila_inicial" y hacia abajo.
# Formato:  "campo": "COLUMNA"
#
# Campos disponibles:
#   "pallet_id", "destination", "packaging", "weight",
#   "line", "material", "description", "quantity",
#   "work_order", "employee_id", "employee_name",
#   "created_at", "printed_at", "status"
# --------------------------------------------------
TABLA = {
    "fila_inicial": 6,
    "columnas": {
        "pallet_id": "A",
        "material": "C",
        "description": "D",
        "quantity": "E",
        "work_order": "F",
        "destination": "G",
    },
}

# Si tu plantilla tiene espacio limitado (por ejemplo, hay un total
# justo debajo de la tabla), pon aquí la última fila disponible.
# Si hay más datos, avisa con un error en lugar de pisar tu formato.
# None = sin límite.
FILA_MAXIMA = None

# Copiar el formato (bordes, fuente, etc.) de la primera fila de la
# tabla a las filas nuevas, para que todas se vean igual.
COPIAR_ESTILO_FILAS = True


# ==================================================
# CONSULTA DE DATOS
# ==================================================

# El orden de estas columnas debe coincidir con CAMPOS.
CONSULTA_ESCANEOS = """
    SELECT
        p.pallet_id,
        p.status,
        p.created_at,
        p.printed_at,
        d.code,
        u.employee_id,
        u.name,
        p.packaging_type,
        p.weight_lb,
        i.line_no,
        i.material,
        m.description,
        i.quantity,
        i.work_order,
        p.qr_data
    FROM pallets p
    JOIN users u
        ON u.id = p.user_id
    JOIN destinations d
        ON d.id = p.destination_id
    LEFT JOIN pallet_items i
        ON i.pallet_record_id = p.id
    LEFT JOIN materials m
        ON m.material_number = i.material
"""

ORDEN_ESCANEOS = "ORDER BY p.id, i.line_no"

CAMPOS = [
    "pallet_id",
    "status",
    "created_at",
    "printed_at",
    "destination",
    "employee_id",
    "employee_name",
    "packaging",
    "weight",
    "line",
    "material",
    "description",
    "quantity",
    "work_order",
    "qr_data",
]


def obtener_escaneos(pallet_ids=None):

    """
    Si pallet_ids es None devuelve todas las tarimas.
    Si recibe una lista, devuelve solo esas, en el mismo orden
    en que fueron escaneadas.
    """

    consulta = CONSULTA_ESCANEOS
    parametros = ()

    if pallet_ids is not None:

        if not pallet_ids:
            return []

        marcas = ",".join("?" * len(pallet_ids))
        consulta += f" WHERE p.pallet_id IN ({marcas}) "
        parametros = tuple(pallet_ids)

    consulta += " " + ORDEN_ESCANEOS

    conexion = conectar()

    try:
        cursor = conexion.cursor()
        cursor.execute(consulta, parametros)
        filas = cursor.fetchall()
    finally:
        conexion.close()

    if pallet_ids is not None:

        posicion = {pid: i for i, pid in enumerate(pallet_ids)}
        filas.sort(key=lambda f: posicion.get(f[0], 10**9))

    return filas


# ==================================================
# 1) LLENAR TU PLANTILLA
# ==================================================

def _a_diccionarios(registros):

    return [dict(zip(CAMPOS, fila)) for fila in registros]


def _escribir(hoja, celda, valor):

    """Escribe en una celda y da un error claro si está combinada."""

    try:
        hoja[celda] = valor
    except AttributeError:
        raise ValueError(
            f"Cell {celda} is part of a merged range. "
            "Use the top-left cell of the merged range."
        )


def generar_packing_desde_plantilla(
    ruta_plantilla,
    ruta_salida,
    registros,
    usuario=None
):

    ruta_plantilla = Path(ruta_plantilla)

    if not ruta_plantilla.exists():
        raise FileNotFoundError(
            f"Template not found: {ruta_plantilla}"
        )

    if Path(ruta_salida).resolve() == ruta_plantilla.resolve():
        raise ValueError(
            "The output file must be different from the template."
        )

    filas = _a_diccionarios(registros)

    wb = load_workbook(ruta_plantilla)

    try:

        if HOJA_PLANTILLA is None:
            hoja = wb.active
        else:
            if HOJA_PLANTILLA not in wb.sheetnames:
                raise ValueError(
                    f"Sheet '{HOJA_PLANTILLA}' not found in the template."
                )
            hoja = wb[HOJA_PLANTILLA]

        usuario = usuario or {}

        # ---------- celdas fijas ----------

        total_cantidad = sum(
            (f["quantity"] or 0) for f in filas
        )

        valores_fijos = {
            "date": datetime.now().strftime("%m/%d/%Y %H:%M"),
            "user_name": usuario.get("name", ""),
            "user_id": usuario.get("employee_id", ""),
            "total_pallets": len({f["pallet_id"] for f in filas}),
            "total_lines": sum(
                1 for f in filas if f["material"] is not None
            ),
            "total_qty": total_cantidad,
            "destination": filas[0]["destination"] if filas else "",
        }

        for campo, celda in CELDAS_FIJAS.items():

            if campo not in valores_fijos:
                raise ValueError(
                    f"Unknown fixed field '{campo}' in CELDAS_FIJAS."
                )

            _escribir(hoja, celda, valores_fijos[campo])

        # ---------- tabla ----------

        fila_inicial = TABLA["fila_inicial"]
        columnas = TABLA["columnas"]

        for campo in columnas:
            if campo not in CAMPOS:
                raise ValueError(
                    f"Unknown table field '{campo}' in TABLA."
                )

        if FILA_MAXIMA is not None:

            ultima = fila_inicial + len(filas) - 1

            if ultima > FILA_MAXIMA:
                raise ValueError(
                    f"The data needs rows up to {ultima}, but the "
                    f"template only allows up to row {FILA_MAXIMA}."
                )

        for i, fila in enumerate(filas):

            numero_fila = fila_inicial + i

            for campo, columna in columnas.items():

                celda = hoja[f"{columna}{numero_fila}"]

                if (
                    COPIAR_ESTILO_FILAS
                    and numero_fila != fila_inicial
                ):
                    origen = hoja[f"{columna}{fila_inicial}"]
                    if origen.has_style:
                        celda._style = copy(origen._style)

                valor = fila[campo]
                celda.value = "" if valor is None else valor

        wb.save(ruta_salida)

    finally:
        wb.close()

    return ruta_salida


# ==================================================
# 2) EXCEL DE QR ESCANEADOS (solo informativo)
# ==================================================

FUENTE = "Arial"
COLOR_ENCABEZADO = "1F4E78"

BORDE = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

ENCABEZADOS_ESCANEOS = [
    "PALLET ID",
    "STATUS",
    "CREATED AT",
    "PRINTED AT",
    "DESTINATION",
    "EMPLOYEE ID",
    "EMPLOYEE NAME",
    "PACKAGING",
    "WEIGHT (LB)",
    "LINE",
    "MATERIAL",
    "DESCRIPTION",
    "QUANTITY",
    "WORK ORDER",
    "QR DATA",
]


def generar_excel_escaneos(ruta_destino, registros=None):

    if registros is None:
        registros = obtener_escaneos()

    wb = Workbook()
    hoja = wb.active
    hoja.title = "SCANS"

    hoja.append(ENCABEZADOS_ESCANEOS)

    for col in range(1, len(ENCABEZADOS_ESCANEOS) + 1):

        celda = hoja.cell(row=1, column=col)
        celda.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
        celda.alignment = Alignment(
            horizontal="center", vertical="center"
        )
        celda.border = BORDE

    for registro in registros:
        hoja.append(["" if v is None else v for v in registro])

    total_filas = hoja.max_row

    for fila in hoja.iter_rows(min_row=2, max_row=total_filas):
        for celda in fila:
            celda.font = Font(name=FUENTE, size=10)
            celda.border = BORDE

    for col in range(1, len(ENCABEZADOS_ESCANEOS) + 1):

        letra = get_column_letter(col)
        mayor = max(
            len(str(c.value)) if c.value is not None else 0
            for c in hoja[letra]
        )
        hoja.column_dimensions[letra].width = min(max(mayor + 3, 12), 60)

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    fila_total = total_filas + 2
    hoja.cell(row=fila_total, column=1, value="TOTAL LINES").font = Font(
        name=FUENTE, bold=True
    )
    hoja.cell(
        row=fila_total,
        column=2,
        value=f"=COUNTA(A2:A{total_filas})"
    ).font = Font(name=FUENTE, bold=True)

    # Solo lectura (sin contraseña: se puede quitar desde Excel)
    hoja.protection.sheet = True
    hoja.protection.autoFilter = False
    hoja.protection.sort = False

    wb.save(ruta_destino)
    return ruta_destino


# ==================================================
# 3) GENERAR LOS DOS ARCHIVOS
# ==================================================

def generar_ambos_excel(
    carpeta,
    pallet_ids=None,
    ruta_plantilla=None,
    usuario=None
):

    carpeta = Path(carpeta)
    sello = datetime.now().strftime("%Y%m%d_%H%M%S")

    ruta_packing = carpeta / f"packing_list_{sello}.xlsx"
    ruta_escaneos = carpeta / f"scanned_qr_{sello}.xlsx"

    registros = obtener_escaneos(pallet_ids)

    generar_packing_desde_plantilla(
        ruta_plantilla or PLANTILLA_PACKING,
        ruta_packing,
        registros,
        usuario
    )

    generar_excel_escaneos(ruta_escaneos, registros)

    return ruta_packing, ruta_escaneos
